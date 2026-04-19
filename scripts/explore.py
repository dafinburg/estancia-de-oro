"""Explorar archivos con más detalle"""
import openpyxl

DOWNLOADS = "C:/Users/HP/Downloads"

# Ver el formato completo de la nota de pedido
print("=" * 60)
print("FORMATO NOTA DE PEDIDO - TODAS LAS FILAS")
print("=" * 60)
wb = openpyxl.load_workbook(f"{DOWNLOADS}/FORMATO NOTA DE PEDIDO (1).xlsx", data_only=True)
ws = wb["pedido d d"]
for i, row in enumerate(ws.iter_rows(values_only=True)):
    print(f"Row {i}: {row}")

# Ver todas las listas de precios
print("\n" + "=" * 60)
print("TABLA LISTAS DE PRECIOS - TODAS LAS FILAS")
print("=" * 60)
wb2 = openpyxl.load_workbook(f"{DOWNLOADS}/Tabla listas de precios_clientes.xlsx", data_only=True)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    print(f"Sheet: {sn}")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        print(f"  Row {i}: {row}")

# Ver productos (listado articulos) - Solo los primeros 50
print("\n" + "=" * 60)
print("LISTADO ARTICULOS - PRIMEROS 50")
print("=" * 60)
wb3 = openpyxl.load_workbook(f"{DOWNLOADS}/LISTADO DE ARTICULOS.xlsx", data_only=True)
ws = wb3["Hoja1"]
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i < 50:
        r = tuple(c for c in row if c is not None)
        if r:
            print(f"Row {i}: {r}")
